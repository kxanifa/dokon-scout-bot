import hashlib
import hmac
import io
import json
import time
from unittest.mock import patch

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.api.deps import get_current_admin, get_current_superadmin
from app.main import app
from app.services.auth import AuthError, validate_telegram_init_data
from app.services.excel import generate_stores_excel
from app.services.sheets import Agent, Store, sheets_service


def create_test_init_data(user_dict: dict, bot_token: str, auth_date: int = None) -> str:
    """Helper to generate a valid Telegram initData string with valid HMAC-SHA256 signature."""
    if auth_date is None:
        auth_date = int(time.time())

    params = {
        "auth_date": str(auth_date),
        "query_id": "AAHdF6IQAAAAAN0XohDhrOrc",
        "user": json.dumps(user_dict, separators=(",", ":")),
    }

    sorted_pairs = [f"{k}={v}" for k, v in sorted(params.items())]
    data_check_string = "\n".join(sorted_pairs)

    secret_key = hmac.new(b"WebAppData", bot_token.encode("utf-8"), hashlib.sha256).digest()
    calc_hash = hmac.new(secret_key, data_check_string.encode("utf-8"), hashlib.sha256).hexdigest()

    params["hash"] = calc_hash
    return "&".join(f"{k}={v}" for k, v in params.items())


def test_init_data_validation():
    bot_token = "1234567890:ABCDefGhIJKlmNoPQRsTUVwxyZ-12345"
    user_dict = {"id": 999111, "first_name": "Admin", "username": "admin_user"}

    # 1. Valid initData
    valid_init_data = create_test_init_data(user_dict, bot_token)
    parsed = validate_telegram_init_data(valid_init_data, bot_token)
    assert parsed["user_obj"]["id"] == 999111

    # 2. Corrupted hash
    bad_init_data = valid_init_data.replace("hash=", "hash=deadbeef")
    with pytest.raises(AuthError):
        validate_telegram_init_data(bad_init_data, bot_token)

    # 3. Expired auth_date (> 24 hours ago)
    expired_date = int(time.time()) - 90000
    expired_init_data = create_test_init_data(user_dict, bot_token, auth_date=expired_date)
    with pytest.raises(AuthError):
        validate_telegram_init_data(expired_init_data, bot_token)

    # 4. Missing hash
    with pytest.raises(AuthError):
        validate_telegram_init_data("auth_date=12345&user={}", bot_token)


def test_excel_export_generation():
    stores = [
        Store(
            id=1,
            name="Test Store 1",
            inn="012345678",  # with leading zero
            phone="+998901234567",
            state="Toshkent shahri",
            district="Yunusobod",
            mahalla="Markaz-4",
            lat=41.31,
            lon=69.24,
            status="faol",
            agent_id=100,
            agent_name="Agent 100",
        ),
        Store(
            id=2,
            name="Test Store 2",
            inn="987654321",
            phone="+998907654321",
            state="Toshkent viloyati",
            district="Chirchiq",
            mahalla="1-mavze",
            lat=41.45,
            lon=69.58,
            status="faol",
            agent_id=100,
            agent_name="Agent 100",
        ),
    ]
    agents = {
        100: Agent(telegram_id=100, name="Agent 100", phone="+998901112233", role="agent"),
    }

    excel_bytes = generate_stores_excel(stores, agents)
    assert len(excel_bytes) > 0

    # Load back with openpyxl to inspect structure
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Do'konlar" in wb.sheetnames
    assert "Hududlar bo'yicha" in wb.sheetnames
    assert "Agentlar bo'yicha" in wb.sheetnames

    ws1 = wb["Do'konlar"]
    assert ws1.cell(row=1, column=1).value == "ID"
    assert ws1.cell(row=2, column=5).value == "Test Store 1"
    # Verify INN formatted as text
    cell_inn = ws1.cell(row=2, column=6)
    assert cell_inn.number_format == "@"

    ws2 = wb["Hududlar bo'yicha"]
    assert ws2.cell(row=1, column=1).value == "Viloyat"


def test_api_photo_security():
    client = TestClient(app)

    mock_admin = Agent(telegram_id=1, name="Admin", role="admin", status="active")
    app.dependency_overrides[get_current_admin] = lambda: mock_admin

    # Mock sheets service with known stores
    known_store = Store(id=1, photo1_id="valid_photo_id_123")
    with patch.object(sheets_service, "get_stores", return_value=[known_store]), patch.object(sheets_service, "get_visits", return_value=[]):
        # 1. Unregistered file ID must return 404 (prevents reading arbitrary Drive files)
        resp = client.get("/api/photo/arbitrary_drive_file_id")
        assert resp.status_code == 404

    app.dependency_overrides.clear()


def test_api_stores_filtering_and_pagination():
    client = TestClient(app)

    mock_admin = Agent(telegram_id=1, name="Admin", role="admin", status="active")
    app.dependency_overrides[get_current_admin] = lambda: mock_admin

    stores = [
        Store(id=1, name="Baraka Market", inn="111111111", state="Toshkent shahri", district="Chilonzor", status="faol"),
        Store(id=2, name="Grand Supermarket", inn="222222222", state="Samarqand viloyati", district="Urgut", status="faol"),
        Store(id=3, name="Baraka Mini", inn="333333333", state="Toshkent shahri", district="Yunusobod", status="faol"),
    ]

    with patch.object(sheets_service, "get_stores", return_value=stores):
        # Filter by region
        r1 = client.get("/api/stores?viloyat=Toshkent shahri")
        assert r1.status_code == 200
        assert r1.json()["total"] == 2

        # Search query
        r2 = client.get("/api/stores?q=baraka")
        assert r2.status_code == 200
        assert r2.json()["total"] == 2

        # Search by INN
        r3 = client.get("/api/stores?q=222222222")
        assert r3.status_code == 200
        assert r3.json()["total"] == 1
        assert r3.json()["stores"][0]["name"] == "Grand Supermarket"

    app.dependency_overrides.clear()


def test_superadmin_role_enforcement():
    client = TestClient(app)

    regular_admin = Agent(telegram_id=2, name="Admin Normal", role="admin", status="active")
    app.dependency_overrides[get_current_admin] = lambda: regular_admin

    # Regular admin tries superadmin endpoint -> 403 Forbidden
    resp = client.post("/api/admins/99")
    assert resp.status_code == 403

    # Superadmin tries -> 200 OK
    super_admin = Agent(telegram_id=1, name="SuperAdmin", role="superadmin", status="active")
    app.dependency_overrides[get_current_admin] = lambda: super_admin
    app.dependency_overrides[get_current_superadmin] = lambda: super_admin

    with patch.object(sheets_service, "set_agent_role", return_value=True):
        resp2 = client.post("/api/admins/99")
        assert resp2.status_code == 200
        assert resp2.json() == {"ok": True}

    app.dependency_overrides.clear()


def test_admin_login_and_bearer_auth():
    client = TestClient(app)
    from app.config import get_settings
    settings = get_settings()

    # 1. Bad secret -> 401
    bad_resp = client.post("/api/auth/login", json={"secret": "wrong_secret_12345"})
    assert bad_resp.status_code == 401

    # 2. Valid secret -> 200 with token
    good_resp = client.post("/api/auth/login", json={"secret": settings.WEBHOOK_SECRET})
    assert good_resp.status_code == 200
    data = good_resp.json()
    assert data["ok"] is True
    token = data["token"]
    assert token

    # 3. Access /api/me with Bearer token
    with patch.object(sheets_service, "get_agent_by_id", return_value=Agent(telegram_id=settings.SUPERADMIN_ID, name="SuperAdmin", role="superadmin", status="active")):
        me_resp = client.get("/api/me", headers={"Authorization": f"Bearer {token}"})
        assert me_resp.status_code == 200
        assert me_resp.json()["role"] == "superadmin"
