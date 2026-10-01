import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.sheets import Agent, Store


def generate_stores_excel(
    stores: list[Store],
    agents_map: dict[int, Agent],
) -> bytes:
    """
    Generate professional multi-sheet Excel report according to SPEC section 9:
    - Sheet 1: 'Do'konlar'
    - Sheet 2: 'Hududlar bo'yicha'
    - Sheet 3: 'Agentlar bo'yicha'
    """
    if len(stores) > 50000:
        raise ValueError("Eksport qatorlari soni 50 000 tadan oshmasligi kerak.")

    wb = Workbook()

    # Colors & Styles
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")

    # ------------------ Sheet 1: Do'konlar ------------------
    ws1 = wb.active
    ws1.title = "Do'konlar"

    headers1 = [
        "ID", "Sana", "Vaqt", "Agent Ismi", "Do'kon Nomi", "INN",
        "Telefon", "Viloyat", "Tuman", "Mahalla", "Lat", "Lon", "Holat"
    ]
    ws1.append(headers1)

    for r_idx, s in enumerate(stores, start=2):
        row = [
            s.id,
            s.date,
            s.time,
            s.agent_name,
            s.name,
            str(s.inn),
            s.phone,
            s.state,
            s.district,
            s.mahalla,
            round(s.lat, 6),
            round(s.lon, 6),
            s.status,
        ]
        ws1.append(row)

        # Format INN as text explicitly in cell
        cell_inn = ws1.cell(row=r_idx, column=6)
        cell_inn.number_format = "@"

        # Zebra striping
        if r_idx % 2 == 0:
            for c_idx in range(1, len(row) + 1):
                ws1.cell(row=r_idx, column=c_idx).fill = zebra_fill

    # Style Header for Sheet 1
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = ws1.dimensions

    # Auto-adjust column widths
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ------------------ Sheet 2: Hududlar bo'yicha ------------------
    ws2 = wb.create_sheet(title="Hududlar bo'yicha")
    headers2 = ["Viloyat", "Tuman", "Mahalla", "Do'konlar Soni"]
    ws2.append(headers2)

    region_counts: dict[tuple[str, str, str], int] = {}
    for s in stores:
        key = (s.state or "Noma'lum", s.district or "Noma'lum", s.mahalla or "Noma'lum")
        region_counts[key] = region_counts.get(key, 0) + 1

    sorted_regions = sorted(region_counts.items(), key=lambda x: (x[0][0], x[0][1], -x[1]))
    for r_idx, (r_key, count) in enumerate(sorted_regions, start=2):
        row = [r_key[0], r_key[1], r_key[2], count]
        ws2.append(row)
        if r_idx % 2 == 0:
            for c_idx in range(1, 5):
                ws2.cell(row=r_idx, column=c_idx).fill = zebra_fill

    for col_idx in range(1, 5):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # ------------------ Sheet 3: Agentlar bo'yicha ------------------
    ws3 = wb.create_sheet(title="Agentlar bo'yicha")
    headers3 = ["Telegram ID", "Agent Ismi", "Telefon", "Rol", "Kiritgan Do'konlari"]
    ws3.append(headers3)

    agent_counts: dict[int, int] = {}
    for s in stores:
        agent_counts[s.agent_id] = agent_counts.get(s.agent_id, 0) + 1

    sorted_agent_items = sorted(
        agents_map.values(),
        key=lambda ag: agent_counts.get(ag.telegram_id, 0),
        reverse=True,
    )

    for r_idx, ag in enumerate(sorted_agent_items, start=2):
        count = agent_counts.get(ag.telegram_id, 0)
        row = [str(ag.telegram_id), ag.name, ag.phone, ag.role, count]
        ws3.append(row)
        if r_idx % 2 == 0:
            for c_idx in range(1, 6):
                ws3.cell(row=r_idx, column=c_idx).fill = zebra_fill

    for col_idx in range(1, 6):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = ws3.dimensions
    for col in ws3.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Output to bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
